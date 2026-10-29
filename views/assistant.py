"""助手对话页：普通对话 + 上传文档本地问答。

两条路径共用 mini_agents.assistant 的同一个 Agent 循环：
- 无知识库：普通聊天（多轮，历史压缩成 context 注入）
- 有知识库：Agent 带 knowledge_search 工具，回答带 [来源] 标记——
  检索（BM25/jieba）纯本地零成本，只有最终回答走 LLM

上传的 pdf/docx 在此抽成文本（pypdf / python-docx），其余文本文件原样读入，
统一落成 .txt 后交给 KnowledgeBase.ingest。
对话日志按会话写入 runs/chat-*/trace.jsonl，可用 view_trace.py 回放。
"""
import asyncio
from datetime import datetime
from io import BytesIO
from pathlib import Path

import streamlit as st

from mini_agents import LLM, TraceLogger
from mini_agents.assistant import chat_reply
from mini_agents.i18n import t
from mini_agents.rag import KnowledgeBase
from views._shared import render_lang_switcher, render_model_settings

UPLOAD_SUFFIXES = ["txt", "md", "markdown", "pdf", "docx", "py", "json", "csv", "log"]


def tempfile_dir():
    """临时目录上下文：入库后片段全在内存，目录即刻释放。"""
    import tempfile
    return tempfile.TemporaryDirectory(prefix="mini_chat_")


def _extract_to_text(upload, folder: Path):
    """把一个上传文件转成 .txt 落进 folder（KnowledgeBase 只认文本文件）。"""
    suffix = Path(upload.name).suffix.lower()
    data = upload.getvalue()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix == ".docx":
        import docx
        document = docx.Document(BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:  # 表格内容也在文档语义里，不能丢
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n\n".join(parts)
    else:
        text = data.decode("utf-8", errors="replace")
    (folder / f"{Path(upload.name).stem}.txt").write_text(text, encoding="utf-8")


def _session_trace() -> TraceLogger:
    """一个会话一份对话日志（追加写），runs/chat-*/trace.jsonl 可回放。"""
    if "chat_trace" not in st.session_state:
        run_dir = Path("runs") / f"chat-{datetime.now():%m%d-%H%M%S}"
        st.session_state["chat_trace"] = TraceLogger(run_dir / "trace.jsonl", echo=False)
    return st.session_state["chat_trace"]


def main():
    render_lang_switcher()
    render_model_settings()

    st.title(t("chat.title"))
    st.caption(t("chat.caption"))

    if "chat_history" not in st.session_state:
        st.session_state["chat_history"] = []

    # ── 文档知识库 ──────────────────────────────────────────────
    uploads = st.file_uploader(t("chat.upload"), accept_multiple_files=True,
                               type=UPLOAD_SUFFIXES)
    col_build, col_clear = st.columns(2)
    if col_build.button(t("chat.ingest_btn"), use_container_width=True,
                        disabled=not uploads):
        with tempfile_dir() as folder:
            for upload in uploads:
                _extract_to_text(upload, folder)
            kb = KnowledgeBase()
            n_chunks = kb.ingest(folder)  # 片段全部复制进内存，临时目录即可释放
        st.session_state["kb"] = kb
        st.session_state["chat_history"].append({
            "role": "assistant",
            "content": t("chat.ingested", n=n_chunks,
                         files="、".join(u.name for u in uploads))})
        st.rerun()
    if col_clear.button(t("chat.clear_kb"), use_container_width=True,
                        disabled="kb" not in st.session_state):
        st.session_state.pop("kb", None)
        st.session_state["chat_history"].append({
            "role": "assistant", "content": t("chat.cleared")})
        st.rerun()

    kb = st.session_state.get("kb")
    if kb:
        st.info(t("chat.kb_active", n=len(kb.chunks), sources="、".join(kb.sources())))

    # ── 历史对话 ────────────────────────────────────────────────
    for m in st.session_state["chat_history"]:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    # ── 输入与回复 ──────────────────────────────────────────────
    prompt = st.chat_input(t("chat.input_hint"))
    if not prompt:
        if "chat_trace" in st.session_state:
            st.caption(t("chat.trace_note", path=st.session_state["chat_trace"].path))
        return

    history = list(st.session_state["chat_history"])
    with st.chat_message("user"):
        st.markdown(prompt)
    with st.chat_message("assistant"):
        reply = None
        with st.spinner(t("chat.thinking")):
            try:
                llm = LLM(trace=_session_trace())
                reply = asyncio.run(chat_reply(prompt, llm, kb=kb, history=history))
            except Exception as exc:  # Key 未配/网关不通都别让页面崩
                reply = t("chat.error", detail=str(exc)[:200])
        st.markdown(reply)
    st.session_state["chat_history"] = history + [
        {"role": "user", "content": prompt},
        {"role": "assistant", "content": reply}]
    st.rerun()


main()
